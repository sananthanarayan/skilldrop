"""Tests for deck-builder's scripts/build_deck.py.

build_deck.py imports python-pptx at module level. When pptx is not installed, a minimal fake
`pptx` package is placed in sys.modules for the duration of load() only, so the pure helpers
(geometry, palettes, brand handling, layout lookup, CLI argument handling) can still be tested.
Anything that really renders a deck is skipped unless python-pptx is installed.
"""
import contextlib
import io
import json
import os
import sys
import types
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

from tests.support import has_module, load, run, workdir, write

SKILL = "deck-builder"
SCRIPT = "build_deck.py"
HAS_PPTX = has_module("pptx")


# ------------------------------------------------------------------ fake pptx


class _FakeEnum:
    def __init__(self, prefix):
        self._prefix = prefix

    def __getattr__(self, name):
        if name.startswith("_"):
            raise AttributeError(name)
        return f"{self._prefix}.{name}"


class _FakeRGBColor(tuple):
    @classmethod
    def from_string(cls, s):
        return cls((int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16)))

    def __str__(self):
        return "%02X%02X%02X" % self


def _fake_pptx_modules():
    def mod(name, **attrs):
        m = types.ModuleType(name)
        m.__dict__.update(attrs)
        return m

    def presentation(*a, **k):
        raise RuntimeError("fake pptx: Presentation is not available in tests")

    return {
        "pptx": mod("pptx", Presentation=presentation),
        "pptx.chart": mod("pptx.chart"),
        "pptx.chart.data": mod("pptx.chart.data", CategoryChartData=object),
        "pptx.dml": mod("pptx.dml"),
        "pptx.dml.color": mod("pptx.dml.color", RGBColor=_FakeRGBColor),
        "pptx.enum": mod("pptx.enum"),
        "pptx.enum.chart": mod("pptx.enum.chart", XL_CHART_TYPE=_FakeEnum("XL_CHART_TYPE"),
                               XL_LEGEND_POSITION=_FakeEnum("XL_LEGEND_POSITION")),
        "pptx.enum.shapes": mod("pptx.enum.shapes", MSO_SHAPE=_FakeEnum("MSO_SHAPE"),
                                PP_PLACEHOLDER=_FakeEnum("PP_PLACEHOLDER")),
        "pptx.enum.text": mod("pptx.enum.text", MSO_ANCHOR=_FakeEnum("MSO_ANCHOR"),
                              PP_ALIGN=_FakeEnum("PP_ALIGN")),
        "pptx.util": mod("pptx.util", Emu=lambda v: int(v),
                         Inches=lambda v: int(v * 914400), Pt=lambda v: int(v * 12700)),
    }


def load_deck():
    """Load build_deck.py with real pptx if present, else with the fake one (restored after)."""
    if HAS_PPTX:
        return load(SKILL, SCRIPT)
    fakes = _fake_pptx_modules()
    saved = {k: sys.modules.get(k) for k in fakes}
    sys.modules.update(fakes)
    try:
        return load(SKILL, SCRIPT)
    finally:
        for k, v in saved.items():
            if v is None:
                sys.modules.pop(k, None)
            else:
                sys.modules[k] = v


EMU_PER_IN = 914400
EMU_PER_PT = 12700


def emu(inches):
    return int(inches * EMU_PER_IN)


class _Stop(Exception):
    pass


# ------------------------------------------------------------------ pure helpers


class GeometryTest(unittest.TestCase):
    def setUp(self):
        self.m = load_deck()

    def canvas(self, w_in, h_in):
        return self.m.Canvas(emu(w_in), emu(h_in))

    def test_design_canvas_maps_one_to_one(self):
        c = self.canvas(13.333, 7.5)
        self.assertEqual(c.x(0.75), emu(0.75))
        self.assertEqual(c.y(1.5), emu(1.5))
        self.assertEqual(int(c.pt(20)), 20 * EMU_PER_PT)

    def test_4_3_canvas_scales_x_and_font_but_not_y(self):
        c = self.canvas(10.0, 7.5)
        self.assertAlmostEqual(c.x(13.333), c.w, delta=2)
        self.assertEqual(c.y(7.5), c.h)
        # font scale follows the tighter axis: 10/13.333 = 0.75 -> 20pt becomes 15pt
        self.assertEqual(int(c.pt(20)), 15 * EMU_PER_PT)

    def test_font_scale_floored_at_70_percent(self):
        c = self.canvas(5.0, 3.0)
        self.assertEqual(c.fr, 0.7)
        self.assertEqual(int(c.pt(10)), 7 * EMU_PER_PT)

    def test_default_region_below_title_bar(self):
        c = self.canvas(13.333, 7.5)
        r = self.m.default_region(c)
        self.assertEqual((r.left, r.top), (emu(0.75), emu(1.5)))
        self.assertEqual(r.width, c.w - 2 * emu(0.75))
        self.assertEqual(r.height, emu(7.5) - emu(1.5) - emu(0.6))

    def test_default_region_without_title_bar_starts_higher(self):
        c = self.canvas(13.333, 7.5)
        r = self.m.default_region(c, below_title_bar=False)
        self.assertEqual(r.top, emu(0.9))
        self.assertEqual(r.height, emu(7.5) - emu(0.9) - emu(0.6))


class PaletteTest(unittest.TestCase):
    def setUp(self):
        self.m = load_deck()
        self.script_dir = Path(self.m.__file__).resolve().parent

    def test_partial_dict_palette_completed_from_defaults(self):
        p = self.m.resolve_palette({"primary": "#112233"}, self.script_dir)
        self.assertEqual(p["primary"], "#112233")
        self.assertEqual(p["accent"], "#FDBB2D")
        self.assertEqual(set(p), {"primary", "secondary", "accent", "background", "text"})

    def test_named_palette_resolved_from_palettes_json(self):
        p = self.m.resolve_palette("forest", self.script_dir)
        self.assertEqual(p["primary"], "#1B5E20")
        self.assertEqual(p["accent"], "#FBC02D")
        self.assertEqual(self.m.WARNINGS, [])

    def test_unknown_palette_name_warns_and_uses_defaults(self):
        with contextlib.redirect_stderr(io.StringIO()) as err:
            p = self.m.resolve_palette("neon-dreams", self.script_dir)
        self.assertEqual(p, self.m.DEFAULT_PALETTE)
        self.assertIn("neon-dreams", self.m.WARNINGS[0])
        self.assertIn("warning:", err.getvalue())

    def test_missing_palettes_file_warns(self):
        with workdir() as d, contextlib.redirect_stderr(io.StringIO()):
            p = self.m.resolve_palette("forest", Path(d) / "scripts")
        self.assertEqual(p, self.m.DEFAULT_PALETTE)
        self.assertIn("palettes file missing", self.m.WARNINGS[0])

    def test_no_palette_uses_defaults(self):
        self.assertEqual(self.m.resolve_palette(None, self.script_dir), self.m.DEFAULT_PALETTE)

    def test_series_colors_order_and_cycle(self):
        pal = {"primary": "#000001", "accent": "#000002", "secondary": "#000003",
               "text": "#000004"}
        colors = self.m.series_colors(pal, 8)
        self.assertEqual(colors[:4], ["#000001", "#000002", "#000003", "#000004"])
        self.assertEqual(colors[6:], ["#000001", "#000002"])

    def test_hex_color_accepts_lowercase_and_hash(self):
        self.assertEqual(str(self.m.hex_color("#1a2a6c")), "1A2A6C")
        self.assertEqual(str(self.m.hex_color("FDBB2D")), "FDBB2D")

    def test_is_dark_by_relative_luminance(self):
        self.assertTrue(self.m.is_dark("#000000"))
        self.assertTrue(self.m.is_dark("#1A2A6C"))
        self.assertFalse(self.m.is_dark("#FFFFFF"))
        self.assertFalse(self.m.is_dark("FDBB2D"))
        self.assertFalse(self.m.is_dark("#808080"))  # luminance 0.502, just over 0.45


class BrandTest(unittest.TestCase):
    def setUp(self):
        self.m = load_deck()

    def quiet(self):
        return contextlib.redirect_stderr(io.StringIO())

    def test_load_brand_inline_object(self):
        brand, base = self.m.load_brand({"colors": {"primary": "#123456"}}, Path("x"))
        self.assertEqual(brand["colors"]["primary"], "#123456")
        self.assertEqual(base, Path("x"))

    def test_load_brand_relative_path_resolves_from_base(self):
        with workdir() as d:
            write(os.path.join(d, "brand", "brand.json"), '{"name": "Fictional Co"}')
            brand, base = self.m.load_brand(os.path.join("brand", "brand.json"), Path(d))
            self.assertEqual(brand, {"name": "Fictional Co"})
            self.assertEqual(base, Path(d) / "brand")

    def test_load_brand_missing_file_warns_and_continues(self):
        with workdir() as d, self.quiet():
            brand, _ = self.m.load_brand("nope.json", Path(d))
        self.assertEqual(brand, {})
        self.assertIn("brand file not found", self.m.WARNINGS[0])

    def test_load_brand_invalid_json_warns_and_continues(self):
        with workdir() as d, self.quiet():
            write(os.path.join(d, "brand.json"), "{not json")
            brand, _ = self.m.load_brand("brand.json", Path(d))
        self.assertEqual(brand, {})
        self.assertIn("not valid JSON", self.m.WARNINGS[0])

    def test_load_brand_expands_home_in_spec_path(self):
        with workdir() as home, workdir() as specdir, self.quiet():
            write(os.path.join(home, "brand.json"), '{"name": "Fictional Co"}')
            with mock.patch.dict(os.environ, {"HOME": home, "USERPROFILE": home}):
                brand, _ = self.m.load_brand("~/brand.json", Path(specdir))
        self.assertEqual(brand, {"name": "Fictional Co"})

    def test_brand_colors_fill_missing_palette_dropping_unknown_and_empty(self):
        spec = {}
        brand = {"colors": {"primary": "#101010", "accent": "", "neon": "#00FF00"}}
        self.m.apply_brand(spec, brand, Path("."), Path("."))
        self.assertEqual(spec["palette"], {"primary": "#101010"})

    def test_spec_palette_wins_over_brand_colors(self):
        spec = {"palette": {"accent": "#ABCDEF"}}
        brand = {"colors": {"primary": "#101010", "accent": "#202020"}}
        self.m.apply_brand(spec, brand, Path("."), Path("."))
        self.assertEqual(spec["palette"], {"primary": "#101010", "accent": "#ABCDEF"})

    def test_named_spec_palette_is_kept_over_brand_colors(self):
        spec = {"palette": "forest"}
        self.m.apply_brand(spec, {"colors": {"primary": "#101010"}}, Path("."), Path("."))
        self.assertEqual(spec["palette"], "forest")

    def test_brand_template_resolved_relative_to_brand_dir(self):
        with workdir() as d:
            spec = {}
            self.m.apply_brand(spec, {"templates": {"pptx": "master.potx"}}, Path(d), Path("."))
            self.assertEqual(Path(spec["template"]), (Path(d) / "master.potx").resolve())

    def test_spec_template_wins_over_brand_template(self):
        spec = {"template": "mine.pptx"}
        self.m.apply_brand(spec, {"template": "theirs.pptx"}, Path("."), Path("."))
        self.assertEqual(spec["template"], "mine.pptx")

    def test_fonts_accept_family_objects_or_plain_strings(self):
        self.m.apply_brand({}, {"fonts": {"heading": {"family": "Fakeface Display"},
                                          "body": "Fakeface Text"}}, Path("."), Path("."))
        self.assertEqual(self.m.BRAND["heading_font"], "Fakeface Display")
        self.assertEqual(self.m.BRAND["body_font"], "Fakeface Text")

    def test_logo_png_and_on_dark_variant_registered(self):
        with workdir() as d:
            write(os.path.join(d, "logo.png"), "png")
            write(os.path.join(d, "logo-white.PNG"), "png")
            self.m.apply_brand({}, {"logo": {"primary": "logo.png", "on_dark": "logo-white.PNG"}},
                               Path(d), Path("."))
            self.assertEqual(self.m.BRAND["logo"], str(Path(d) / "logo.png"))
            self.assertEqual(self.m.BRAND["logo_on_dark"], str(Path(d) / "logo-white.PNG"))
        self.assertEqual(self.m.WARNINGS, [])

    def test_svg_logo_warns_and_is_not_used(self):
        with workdir() as d, self.quiet():
            write(os.path.join(d, "logo.svg"), "<svg/>")
            self.m.apply_brand({}, {"logo": "logo.svg"}, Path(d), Path("."))
        self.assertIsNone(self.m.BRAND["logo"])
        self.assertIn("PNG", self.m.WARNINGS[0])

    def test_missing_logo_warns(self):
        with workdir() as d, self.quiet():
            self.m.apply_brand({}, {"logo": "absent.png"}, Path(d), Path("."))
        self.assertIsNone(self.m.BRAND["logo"])
        self.assertIn("logo not found", self.m.WARNINGS[0])


def _layout(name, n_shapes=1):
    el = SimpleNamespace(findall=lambda _q, n=n_shapes: ["sp"] * n)
    return SimpleNamespace(name=name, placeholders=SimpleNamespace(_element=el))


def _shape(kind):
    return SimpleNamespace(placeholder_format=SimpleNamespace(type=kind))


class TemplateHelpersTest(unittest.TestCase):
    def setUp(self):
        self.m = load_deck()

    def prs(self, *layouts):
        return SimpleNamespace(slide_layouts=list(layouts))

    def test_find_layout_prefers_exact_name_over_substring(self):
        prs = self.prs(_layout("Title and Content"), _layout("Title"))
        self.assertEqual(self.m.find_layout(prs, "Title").name, "Title")

    def test_find_layout_case_insensitive_then_substring(self):
        prs = self.prs(_layout("Section Header"), _layout("Title and Content"))
        self.assertEqual(self.m.find_layout(prs, "section header").name, "Section Header")
        self.assertEqual(self.m.find_layout(prs, "content").name, "Title and Content")
        self.assertEqual(self.m.WARNINGS, [])

    def test_find_layout_by_index_and_out_of_range(self):
        prs = self.prs(_layout("A"), _layout("B"))
        self.assertEqual(self.m.find_layout(prs, 1).name, "B")
        with contextlib.redirect_stderr(io.StringIO()):
            self.assertIsNone(self.m.find_layout(prs, 5))
        self.assertIn("out of range (2 layouts)", self.m.WARNINGS[0])

    def test_find_layout_unknown_name_warns(self):
        with contextlib.redirect_stderr(io.StringIO()):
            self.assertIsNone(self.m.find_layout(self.prs(_layout("A")), "Agenda"))
        self.assertIn("'Agenda' not found", self.m.WARNINGS[0])

    def test_list_layouts_prints_index_and_name(self):
        prs = self.prs(_layout("Title Slide"), _layout("Blank"))
        self.assertEqual(self.m.list_layouts(prs), "  [0] Title Slide\n  [1] Blank")

    def test_blank_layout_picks_first_without_shapes(self):
        prs = self.prs(_layout("Title"), _layout("Empty", 0), _layout("Blank"))
        self.assertEqual(self.m.blank_layout(prs).name, "Empty")

    def test_blank_layout_falls_back_to_name_then_last(self):
        prs = self.prs(_layout("Title"), _layout(" Blank Slide "))
        self.assertEqual(self.m.blank_layout(prs).name, " Blank Slide ")
        prs = self.prs(_layout("Title"), _layout("Content"))
        self.assertEqual(self.m.blank_layout(prs).name, "Content")

    def test_placeholders_skip_footer_and_filter_by_kind(self):
        P = self.m.PP_PLACEHOLDER
        title, body, footer = _shape(P.TITLE), _shape(P.BODY), _shape(P.FOOTER)
        slide = SimpleNamespace(placeholders=[title, body, footer, SimpleNamespace()])
        self.assertEqual(self.m.placeholders(slide, self.m.TITLE_PH), [title])
        everything = self.m.placeholders(slide)
        self.assertNotIn(footer, everything)
        self.assertIn(body, everything)

    def test_layout_numbers_slides_detects_slide_number_placeholder(self):
        P = self.m.PP_PLACEHOLDER
        numbered = SimpleNamespace(placeholders=[_shape(P.TITLE), _shape(P.SLIDE_NUMBER)])
        plain = SimpleNamespace(placeholders=[_shape(P.TITLE)])
        self.assertTrue(self.m.layout_numbers_slides(numbered))
        self.assertFalse(self.m.layout_numbers_slides(plain))

    def test_potx_repacked_as_pptx_content_type(self):
        ct = ('<Types><Override PartName="/ppt/presentation.xml" ContentType="application/'
              'vnd.openxmlformats-officedocument.presentationml.template.main+xml"/></Types>')
        with workdir() as d:
            src = Path(d) / "brand.potx"
            with zipfile.ZipFile(src, "w") as z:
                z.writestr("[Content_Types].xml", ct)
                z.writestr("ppt/slides/slide1.xml", "<sld/>")
            out_dir = Path(d) / "out"
            out_dir.mkdir()
            dst = self.m._potx_to_pptx(src, out_dir)
            self.assertEqual(dst.name, "brand.pptx")
            with zipfile.ZipFile(dst) as z:
                types_xml = z.read("[Content_Types].xml").decode("utf-8")
                self.assertEqual(z.read("ppt/slides/slide1.xml"), b"<sld/>")
        self.assertIn("presentationml.presentation.main+xml", types_xml)
        self.assertNotIn("template.main", types_xml)


class DrawingDecisionsTest(unittest.TestCase):
    """The size and layout choices the draw helpers make, with the pptx calls stubbed out."""

    def setUp(self):
        self.m = load_deck()
        self.c = self.m.Canvas(emu(13.333), emu(7.5))
        self.region = self.m.default_region(self.c)
        self.pal = dict(self.m.DEFAULT_PALETTE)
        self.texts = []
        self.m.add_textbox = lambda slide, text, **kw: self.texts.append((text, kw))

    def test_big_number_size_steps_with_length(self):
        for number, pt in (("$24M", 140), ("1,234,567", 100), ("12,345,678,901", 72)):
            self.texts.clear()
            self.m.draw_big_number(None, self.c, self.pal, {"number": number}, self.region)
            self.assertEqual(self.texts[0][0], number)
            self.assertEqual(int(self.texts[0][1]["size"]), pt * EMU_PER_PT, number)

    def test_big_number_caption_and_missing_number(self):
        self.m.draw_big_number(None, self.c, self.pal, {"caption": "ARR at risk"}, self.region)
        self.assertEqual([t for t, _ in self.texts], ["—", "ARR at risk"])

    def test_table_over_12_rows_warns_and_height_is_capped(self):
        placed = []
        self.m.place_table = lambda slide, cols, rows, region, c, pal: placed.append(region)
        rows = [[str(i), "x"] for i in range(13)]
        with contextlib.redirect_stderr(io.StringIO()):
            self.m.draw_table(None, self.c, self.pal, {"columns": ["A", "B"], "rows": rows},
                              self.region, {})
        self.assertIn("13 data rows", self.m.WARNINGS[0])
        self.assertEqual(placed[0].height, min(self.region.height, emu(0.45) * 14))

    def test_short_table_height_is_row_count_times_045in(self):
        placed = []
        self.m.place_table = lambda slide, cols, rows, region, c, pal: placed.append(region)
        self.m.draw_table(None, self.c, self.pal, {"columns": ["A"], "rows": [["1"], ["2"]]},
                          self.region, {})
        self.assertEqual(placed[0].height, emu(0.45) * 3)
        self.assertEqual(self.m.WARNINGS, [])

    def test_empty_table_is_skipped_with_warning(self):
        self.m.place_table = lambda *a: self.fail("table drawn")
        with contextlib.redirect_stderr(io.StringIO()):
            self.m.draw_table(None, self.c, self.pal, {}, self.region, {})
        self.assertIn("no columns or rows", self.m.WARNINGS[0])

    def chart_slide(self, captured):
        def add_chart(kind, *a):
            captured["kind"] = kind
            captured["data"] = a[-1]
            raise _Stop

        return SimpleNamespace(shapes=SimpleNamespace(add_chart=add_chart))

    def fake_chart_data(self):
        class Data:
            def __init__(self):
                self.series = []

            def add_series(self, name, values):
                self.series.append((name, values))

        self.m.CategoryChartData = Data

    def test_chart_without_series_is_skipped_with_warning(self):
        slide = self.chart_slide({})
        with contextlib.redirect_stderr(io.StringIO()):
            self.m.place_chart(slide, {"categories": ["Q1"]}, self.region, self.c, self.pal)
        self.assertIn("no categories or no series", self.m.WARNINGS[0])

    def test_unknown_chart_type_falls_back_to_column(self):
        self.fake_chart_data()
        got = {}
        spec = {"type": "Radar", "categories": ["Q1"], "series": [{"name": "A", "values": [1]}]}
        with contextlib.redirect_stderr(io.StringIO()), self.assertRaises(_Stop):
            self.m.place_chart(self.chart_slide(got), spec, self.region, self.c, self.pal)
        self.assertEqual(got["kind"], self.m.CHART_TYPES["column"])
        self.assertIn("'radar'", self.m.WARNINGS[0])

    def test_series_padded_or_trimmed_to_category_count(self):
        self.fake_chart_data()
        got = {}
        spec = {"type": "line", "categories": ["Q1", "Q2", "Q3"],
                "series": [{"name": "Short", "values": [1]},
                           {"name": "Long", "values": [1, 2, 3, 4]},
                           {"values": [5, 6, 7]}]}
        with contextlib.redirect_stderr(io.StringIO()), self.assertRaises(_Stop):
            self.m.place_chart(self.chart_slide(got), spec, self.region, self.c, self.pal)
        self.assertEqual(got["kind"], self.m.CHART_TYPES["line"])
        self.assertEqual(got["data"].series, [("Short", [1, None, None]),
                                              ("Long", [1, 2, 3]),
                                              ("Series", [5, 6, 7])])
        self.assertEqual(len(self.m.WARNINGS), 2)

    def test_title_slide_footer_joins_presenter_and_date(self):
        self.m.set_background = lambda *a: None
        self.m.add_rect = lambda *a, **k: None
        spec = {"title": "Plan", "presenter": "A. Presenter", "date": "Q3"}
        self.m.render_title(None, self.c, self.pal, spec, {})
        self.assertEqual([t for t, _ in self.texts], ["Plan", "A. Presenter  ·  Q3"])

    def test_logo_variant_follows_slide_darkness(self):
        pics = []
        slide = SimpleNamespace(shapes=SimpleNamespace(
            add_picture=lambda path, *a, **k: pics.append(path) or SimpleNamespace(width=10)))
        self.m.BRAND.update(logo="light.png", logo_on_dark="dark.png")
        for logical in ("title", "content", "section", "closing"):
            self.m.place_logo(slide, self.c, self.pal, logical)
        self.assertEqual(pics, ["light.png", "light.png", "dark.png", "dark.png"])

    def test_no_logo_places_nothing(self):
        slide = SimpleNamespace(shapes=SimpleNamespace(add_picture=lambda *a, **k: self.fail()))
        self.m.place_logo(slide, self.c, self.pal, "title")


class MainArgsTest(unittest.TestCase):
    """main() and build() up to the point where pptx would be needed."""

    def setUp(self):
        self.m = load_deck()

    def call_main(self, *argv):
        err = io.StringIO()
        with contextlib.redirect_stderr(err), contextlib.redirect_stdout(io.StringIO()):
            try:
                rc = self.m.main(list(argv))
            except SystemExit as e:
                rc = e.code
        return rc, err.getvalue()

    def test_no_spec_is_a_usage_error(self):
        rc, err = self.call_main()
        self.assertEqual(rc, 2)
        self.assertIn("spec argument is required", err)

    def test_missing_spec_file_exits_1(self):
        with workdir() as d:
            rc, err = self.call_main(os.path.join(d, "nope.json"))
        self.assertEqual(rc, 1)
        self.assertIn("error: spec file not found", err)

    def test_list_layouts_missing_template_exits_1(self):
        with workdir() as d:
            rc, err = self.call_main("--list-layouts", os.path.join(d, "nope.potx"))
        self.assertEqual(rc, 1)
        self.assertIn("error: template not found", err)

    def test_spec_without_slides_exits_1(self):
        with workdir() as d:
            spec = write(os.path.join(d, "deck.json"), '{"title": "T", "slides": []}')
            rc, err = self.call_main(spec)
        self.assertEqual(rc, 1)
        self.assertIn("error: spec has no slides", err)

    def test_malformed_spec_json_fails_with_clear_error(self):
        with workdir() as d:
            spec = write(os.path.join(d, "deck.json"), '{"slides": [}')
            rc, err = self.call_main(spec)
        self.assertEqual(rc, 1)
        self.assertIn("error:", err)

    def test_output_defaults_next_to_spec_and_flags_pass_through(self):
        calls = []
        self.m.build = lambda spec, out, script_dir, **kw: calls.append((spec, out, kw))
        with workdir() as d:
            spec = write(os.path.join(d, "deck.json"), "{}")
            self.assertEqual(self.call_main(spec, "--strict")[0], 0)
            self.assertEqual(calls[0][1], Path(spec).resolve().with_suffix(".pptx"))
            self.assertTrue(calls[0][2]["strict"])
            self.assertIsNone(calls[0][2]["brand_ref"])
            out = os.path.join(d, "out", "x.pptx")
            self.call_main(spec, "-o", out, "--brand", os.path.join(d, "b.json"))
            self.assertEqual(calls[1][1], Path(out).resolve())
            self.assertFalse(calls[1][2]["strict"])
            self.assertEqual(Path(calls[1][2]["brand_ref"]), Path(d, "b.json").resolve())

    def test_cli_brand_overrides_spec_brand_and_spec_palette_still_wins(self):
        seen = {}

        def stop(spec, spec_dir, tmpdir):
            seen.update(spec)
            raise _Stop

        self.m.open_presentation = stop
        with workdir() as d:
            write(os.path.join(d, "spec-brand.json"), json.dumps({"colors": {"primary": "#111111"}}))
            cli = write(os.path.join(d, "cli-brand.json"),
                        json.dumps({"colors": {"primary": "#222222", "accent": "#333333"}}))
            spec = write(os.path.join(d, "deck.json"), json.dumps({
                "brand": "spec-brand.json", "palette": {"accent": "#444444"},
                "slides": [{"layout": "content", "title": "x"}]}))
            with self.assertRaises(_Stop):
                self.m.build(Path(spec), Path(d, "deck.pptx"), Path(d), strict=False,
                             brand_ref=cli)
        self.assertEqual(seen["palette"], {"primary": "#222222", "accent": "#444444"})


# ------------------------------------------------------------------ CLI without pptx


@unittest.skipIf(HAS_PPTX, "python-pptx is installed")
class CliWithoutPptxTest(unittest.TestCase):
    def test_build_fails_non_zero_naming_pptx(self):
        with workdir() as d:
            spec = write(os.path.join(d, "deck.json"), '{"slides": [{"layout": "title"}]}')
            code, out, err = run(SKILL, SCRIPT, spec, cwd=d)
            self.assertFalse(os.path.exists(os.path.join(d, "deck.pptx")))
        self.assertNotEqual(code, 0)
        self.assertIn("pptx", err)

    def test_missing_pptx_gives_install_hint_not_traceback(self):
        code, out, err = run(SKILL, SCRIPT, "--help")
        self.assertNotIn("Traceback", err)
        self.assertIn("pip install", err)


# ------------------------------------------------------------------ real builds


@unittest.skipUnless(HAS_PPTX, "python-pptx not installed")
class BuildDeckTest(unittest.TestCase):
    SPEC = {
        "title": "Fictional Deck", "aspect": "4:3", "palette": "forest",
        "slides": [
            {"layout": "title", "title": "Plan", "subtitle": "Q3", "notes": "Open."},
            {"layout": "content", "title": "Why", "bullets": ["a", "b"], "notes": "Why."},
            {"layout": "chart", "title": "Growth", "source": "Fictional data",
             "chart": {"type": "column", "categories": ["Q1", "Q2"],
                       "series": [{"name": "ARR", "values": [1, 2]}]}},
            {"layout": "table", "title": "T", "columns": ["A", "B"], "rows": [["1", "2"]]},
            {"layout": "big_number", "title": "N", "number": "$24M", "caption": "c"},
            {"layout": "closing", "title": "Thanks"},
        ],
    }

    def build(self, d, spec, *args):
        path = write(os.path.join(d, "deck.json"), json.dumps(spec))
        out = os.path.join(d, "out", "deck.pptx")
        code, stdout, err = run(SKILL, SCRIPT, path, "-o", out, *args, cwd=d)
        return code, stdout, err, out

    def test_builds_every_slide_at_requested_aspect_with_notes(self):
        from pptx import Presentation
        with workdir() as d:
            code, stdout, err, out = self.build(d, self.SPEC)
            self.assertEqual(code, 0, err)
            prs = Presentation(out)
            self.assertEqual(len(prs.slides), 6)
            self.assertEqual(prs.slide_width, emu(10.0))
            self.assertEqual(prs.slides[0].notes_slide.notes_text_frame.text, "Open.")
        self.assertIn("6 slides", stdout)

    def test_missing_image_warns_and_draws_placeholder(self):
        spec = {"slides": [{"layout": "image", "title": "D", "image": "missing.png"}]}
        with workdir() as d:
            code, stdout, err, out = self.build(d, spec)
            self.assertEqual(code, 0, err)
            self.assertTrue(os.path.exists(out))
        self.assertIn("warning: image not found", err)

    def test_strict_fails_on_missing_image(self):
        spec = {"slides": [{"layout": "image", "title": "D", "image": "missing.png"}]}
        with workdir() as d:
            code, stdout, err, out = self.build(d, spec, "--strict")
        self.assertEqual(code, 1)
        self.assertIn("error: image not found", err)

    def test_unknown_layout_and_aspect_warn_but_build(self):
        spec = {"aspect": "21:9", "slides": [{"layout": "timeline", "title": "x"}]}
        with workdir() as d:
            code, stdout, err, out = self.build(d, spec)
        self.assertEqual(code, 0, err)
        self.assertIn("unknown aspect", err)
        self.assertIn("unknown layout 'timeline'", err)


if __name__ == "__main__":
    unittest.main()
